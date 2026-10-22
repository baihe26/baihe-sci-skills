"""Read saved PPTX objects against explicit goals; never infer scientific or visual acceptance."""
from paths import Path, read, sha, save
from core import GateError, source, writable
from scientific_gate import indexed, resolve_file
import re


def check_profile(contract, base):
    if contract.get('schema_version') != 'style-profile-0.1':
        raise GateError('STYLE_PROFILE_VERSION_UNSUPPORTED')
    if not all(contract.get(k) for k in ('profile_id', 'version', 'scope')):
        raise GateError('STYLE_PROFILE_IDENTITY_REQUIRED')
    ref = contract['source']; source(resolve_file(base, ref['path']), ref['sha256'])
    for row in contract.get('overrides', []):
        if not all(row.get(k) for k in ('rule', 'scope', 'reason', 'authorization_source')):
            raise GateError('STYLE_OVERRIDE_SCOPE_AND_AUTHORIZATION_REQUIRED')
        ref = row['authorization_source']; source(resolve_file(base, ref['path']), ref['sha256'])
    return dict(status='PROFILE_SOURCE_HASH_BOUND_NOT_STYLE_COMPLIANCE',
                profile_id=contract['profile_id'], version=contract['version'], scope=contract['scope'],
                overrides=contract.get('overrides', []), author_acceptance='NOT_INHERITED')


KINDS = {'text_equals', 'fill_equals', 'native_type', 'connector_targets', 'min_font_pt', 'edit_action'}


def check_requirements(contract, graph=None):
    if contract.get('schema_version') != 'delivery-requirements-0.1':
        raise GateError('DELIVERY_REQUIREMENTS_VERSION_UNSUPPORTED')
    rows = indexed(contract['observations'])
    if not rows:
        raise GateError('DELIVERY_OBSERVATIONS_REQUIRED')
    objects = indexed([*graph['objects'], *graph['text_objects'], *graph['arrow_objects'], *graph['groups']]) if graph else None
    for row in rows.values():
        if row.get('kind') not in KINDS or not row.get('object_id') or type(row.get('slide_index', 1)) is not int or row.get('slide_index', 1) < 1:
            raise GateError('DELIVERY_OBSERVATION_INVALID')
        if 'expected' not in row:
            raise GateError('DELIVERY_EXPECTATION_REQUIRED')
        if row['kind']=='fill_equals' and (not isinstance(row['expected'],str) or not re.fullmatch(r'#[0-9A-Fa-f]{6}',row['expected'])):
            raise GateError('DELIVERY_COLOR_INVALID')
        if row['kind']=='edit_action' and not row['expected']:
            raise GateError('DELIVERY_EDIT_ACTION_REQUIRED')
        if row['kind'] == 'min_font_pt' and (type(row['expected']) not in (int, float) or not 0 < row['expected'] <= 300):
            raise GateError('DELIVERY_FONT_INVALID')
        if objects is not None:
            if row['object_id'] not in objects:
                raise GateError('DELIVERY_OBJECT_UNRESOLVED')
            obj = objects[row['object_id']]
            if row['kind'] == 'text_equals' and row['expected'] != obj.get('geometry', {}).get('data', {}).get('text'):
                raise GateError('DELIVERY_TEXT_EXPECTATION_DIFFERS_FROM_APPROVED_SCENE')
            if row['kind']=='fill_equals' and row['expected'].upper()!=(graph['style_tokens'].get(obj.get('style_token'),{}).get('fill') or '').upper():
                raise GateError('DELIVERY_COLOR_EXPECTATION_DIFFERS_FROM_APPROVED_SCENE')
    return list(rows.values())


def verify_delivery(contract_path, artifact, out=None):
    from pptx import Presentation
    from pptx.oxml.ns import qn
    contract_path, artifact = Path(contract_path).resolve(), Path(artifact).resolve()
    contract = read(contract_path); rows = check_requirements(contract)
    bound_artifact = source(artifact); contract_ref = source(contract_path)
    prs = Presentation(str(artifact)); pages = []
    def walk(shapes):
        for s in shapes:
            yield s
            if s.shape_type == 6:
                yield from walk(s.shapes)
    for slide in prs.slides:
        objects = {}
        for s in walk(slide.shapes):
            objects.setdefault(s.name, []).append(s)
        pages.append(objects)
    findings = []
    for row in rows:
        page = row.get('slide_index', 1); kind = row['kind']
        matches = pages[page-1].get(row['object_id'], []) if page <= len(pages) else []
        observed = None; state = 'FAIL'
        if len(matches) == 1:
            s = matches[0]
            if kind == 'edit_action':
                state = 'PENDING_NATIVE_EDIT_TEST'
            elif kind == 'text_equals':
                observed = s.text if s.has_text_frame else None
            elif kind == 'fill_equals':
                try: observed = '#' + str(s.fill.fore_color.rgb)
                except (AttributeError, TypeError, ValueError): pass
            elif kind == 'native_type':
                observed = ('picture' if s.shape_type == 13 else 'group' if s.shape_type == 6 else
                    'table' if s.has_table else 'connector' if s._element.tag == qn('p:cxnSp') else
                    'text' if s.has_text_frame and s.text else 'geometry')
            elif kind == 'connector_targets':
                endpoints = []
                by_id = {v.shape_id: v.name for values in pages[page-1].values() for v in values}
                for tag in ('stCxn', 'endCxn'):
                    nodes = s._element.findall('.//' + qn('a:' + tag))
                    endpoints.append(by_id.get(int(nodes[0].get('id'))) if len(nodes) == 1 else None)
                observed = dict(begin=endpoints[0], end=endpoints[1])
            elif kind == 'min_font_pt':
                sizes = [r.font.size.pt if r.font.size is not None else None for p in s.text_frame.paragraphs for r in p.runs if r.text] if s.has_text_frame else []
                if sizes and all(v is not None for v in sizes):
                    observed = min(sizes); state = 'PASS' if observed >= row['expected'] else 'FAIL'
                else:
                    state = 'PENDING_INHERITED_FONT_READBACK'
            if kind not in ('edit_action', 'min_font_pt'):
                expected = row['expected']
                state = 'PASS' if (observed.upper()==expected.upper() if kind=='fill_equals' and observed else observed==expected) else 'FAIL'
        findings.append(dict(id=row['id'], object_id=row['object_id'], slide_index=page, kind=kind,
            status=state, expected=row.get('expected'), observed=observed, matching_objects=len(matches)))
    source(artifact, bound_artifact['sha256']); source(contract_path, contract_ref['sha256'])
    result = dict(status='FAIL' if any(r['status']=='FAIL' for r in findings) else
        'PENDING_NATIVE_CHECKS' if any(r['status']!='PASS' for r in findings) else 'PASS_SAVED_OBJECT_SCOPE_ONLY',
        artifact=bound_artifact, requirements=contract_ref, observations=findings,
        native_application_reopen='NOT_PERFORMED', visual='PENDING', scientific_truth='NOT_ASSESSED', author_acceptance='PENDING')
    if out is not None:
        out = writable(out)
        if out.exists(): raise FileExistsError(out)
        save(out/'最终对象核对.json', result)
    return result
