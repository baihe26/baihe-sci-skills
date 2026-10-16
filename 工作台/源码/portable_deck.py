"""Optional multi-page assembly of already approved, equal-size SceneGraphs.

Reuses the existing portable scene renderer. Does not design slides or replace
Live Studio's media/content/V03 workflow.
"""
from paths import Path, read, save, sha
from core import writable, GateError, source
from scientific_gate import indexed, resolve_file
from preflight import inspect_pack
from portable_ppt import draw_scene, tag_scene, validate_supported
from provenance import record_trace, verify_trace
from processing_ledger import from_scene, add_unresolved


def scene_inputs(manifest_path):
    manifest_path = Path(manifest_path).resolve()
    manifest = read(manifest_path)
    if manifest.get('schema_version') != 'scene-deck-0.1' or not 1 <= len(manifest.get('slides', [])) <= 100:
        raise GateError('SCENE_DECK_MANIFEST_INVALID')
    indexed(manifest['slides'])
    rows = []
    for row in manifest['slides']:
        path = resolve_file(manifest_path.parent, row['scene_path'])
        if not row.get('expected_scene_sha256') or source(path)['sha256'] != row['expected_scene_sha256']:
            raise GateError('SCENE_DECK_SOURCE_HASH_CONFLICT')
        rows.append(dict(row, scene_path=path))
    return rows


def produce_deck(manifest_path, out):
    from pptx import Presentation
    from pptx.util import Pt
    from pptx_tags import presentation_locations
    from pack_files import fingerprint
    out = writable(out)
    if out.exists(): raise FileExistsError(out)
    manifest_source = source(manifest_path)
    inputs = scene_inputs(manifest_path); plans = []; size = None
    for item in inputs:
        checked = inspect_pack(item['scene_path']); graph = read(item['scene_path']); validate_supported(graph)
        scale = graph['geometry']['coordinate_transform']['scale_pt_per_px']
        physical = tuple(round(graph['canvas'][k] * scale, 8) for k in ('width', 'height'))
        if size is not None and physical != size: raise GateError('MIXED_PAGE_SIZE_REQUIRES_APPROVED_LAYOUT')
        size = physical; plans.append((item, graph, checked))
    out.mkdir(parents=True); prs = Presentation(); prs.slide_width = Pt(size[0]); prs.slide_height = Pt(size[1])
    generated = []
    for item, graph, checked in plans:
        slide, shapes, connectors = draw_scene(prs, graph, checked, item['scene_path'])
        contract = tag_scene(slide, shapes, item['scene_path'], checked, item['id'])
        generated.append((item, checked, graph, shapes, contract, connectors))
    for item, _, checked in plans:
        if fingerprint(item['scene_path']) != checked['input_fingerprint']: raise GateError('INPUT_CHANGED_DURING_PRODUCTION')
    source(manifest_path, manifest_source['sha256'])
    target = out / '可编辑多页图件.pptx'; prs.save(target)
    import zipfile,hashlib
    with zipfile.ZipFile(target) as z:
        media_hashes={hashlib.sha256(z.read(name)).hexdigest() for name in z.namelist() if name.startswith('ppt/media/')}
    if any(a['source_sha256'] not in media_hashes for _,_,checked in plans for a in checked['retained_assets']['assets']):
        raise GateError('SOURCE_IMAGE_BYTES_CHANGED')
    loaded = Presentation(target)
    if len(loaded.slides) != len(plans): raise GateError('SAVED_SLIDE_COUNT_CONFLICT')
    def walk(shapes):
        for shape in shapes:
            yield shape
            if shape.shape_type==6:yield from walk(shape.shapes)
    for number,(_,graph,_) in enumerate(plans):
        actual={shape.name:shape for shape in walk(loaded.slides[number].shapes)}
        for obj in graph['text_objects']:
            if actual[obj['id']].text!=obj['geometry']['data']['text']:raise GateError('SAVED_PAGE_TEXT_CONFLICT')
    locations = presentation_locations(target); reports = []
    for number, (item, checked, graph, shapes, contract, connectors) in enumerate(generated, 1):
        actual = [r for r in locations if r['slide_id'] == item['id']]
        if {r['object_id'] for r in actual} != set(shapes): raise GateError('SAVED_SLIDE_OBJECT_SET_CONFLICT')
        page_out = out / '逐页记录' / ('页_%03d' % number)
        record_trace(page_out, contract, target, [dict(slide_id=item['id'], object_id=r['object_id'], native_id=r['native_id']) for r in actual])
        save(page_out / '生产前检查.json', checked)
        ledger = from_scene(graph)
        for row in ledger['elements']:
            if row['id'] in connectors: row.update(treatment='native_bound_connector', editability='NATIVE_RECTANGLE_ENDPOINTS')
        for asset in checked['retained_assets']['assets']:
            ledger['required_ids'].append(asset['id'])
            ledger['elements'].append(dict(id=asset['id'], source_ref=asset['source_ref'], output_ids=[asset['id']],
                                          treatment='retained_pixels', editability='MOVE_SCALE_REPLACE', content_kind=asset['content_kind'],
                                          source_sha256=asset['source_sha256'], crop_box_normalized=asset['crop_box_normalized'],
                                          uncertain=False, review_status='PENDING'))
        add_unresolved(ledger, checked); save(page_out / '元素处理清单.json', ledger)
        verified = verify_trace(page_out / '来源与引用链.json')
        reports.append(dict(slide_id=item['id'], trace=source(page_out / '来源与引用链.json'),
                            verified=verified, unresolved_ids=checked['reconstruction'].get('uncertain_ids', [])))
    report = dict(status='MULTIPAGE_GENERATED_PENDING_NATIVE_VISUAL_REVIEW', manifest=manifest_source, file=source(target),
                  pages=reports, generated_without_office=True, default_studio_replaced=False,
                  native_reopen='NOT_PERFORMED_BY_THIS_BACKEND', scientific_review='PENDING', author_acceptance='PENDING')
    save(out / '生产验收.json', report)
    if any(p['unresolved_ids'] for p in reports):
        (out / '待核对事项.md').write_text('# 待核对事项\n\n' + ''.join('- ' + p['slide_id'] + '：' + '、'.join(p['unresolved_ids']) + '\n' for p in reports if p['unresolved_ids']) + '\n逐页来源和未重建区域见逐页记录；这不是作者接受或正式可用证明。\n', encoding='utf-8')
    return report


if __name__ == '__main__':
    import sys, json
    print(json.dumps(produce_deck(*sys.argv[1:]), ensure_ascii=False))
