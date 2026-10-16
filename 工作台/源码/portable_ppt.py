"""Optional python-pptx backend. Default COM pipeline remains separate and unchanged."""
from paths import Path, read, save, sha, digest
from core import writable, GateError, source
from preflight import inspect_pack
from native_connectors import rectangle_connector_plan
from processing_ledger import from_scene, validate_ledger, add_unresolved
import json, zipfile


def produce(scene_path,out,source_image=None):
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    scene_path=Path(scene_path).resolve();out=writable(out)
    if out.exists():raise FileExistsError(out)
    checked=inspect_pack(scene_path);g=read(scene_path)
    validate_supported(g)
    scale=g['geometry']['coordinate_transform']['scale_pt_per_px']
    out.mkdir(parents=True)
    prs=Presentation();prs.slide_width=Pt(g['canvas']['width']*scale);prs.slide_height=Pt(g['canvas']['height']*scale)
    slide,shapes,plans=draw_scene(prs,g,checked,scene_path)
    contract=tag_scene(slide,shapes,scene_path,checked,'SCENE')
    from provenance import record_trace,verify_trace
    target=out/'可编辑示意图.pptx';prs.save(str(target))
    # Reload via the same open library, distinct from an Office render/reopen.
    loaded=Presentation(str(target))
    def walk(collection):
        for shape in collection:
            yield shape
            if shape.shape_type==6:yield from walk(shape.shapes)
    loaded_shapes={s.name:s for s in walk(loaded.slides[0].shapes)}
    if set(loaded_shapes)!=set(shapes):raise GateError('PORTABLE_SAVED_ID_SET_CONFLICT')
    for obj in g['text_objects']:
        if loaded_shapes[obj['id']].text!=obj['geometry']['data']['text']:raise GateError('PORTABLE_SAVED_TEXT_CONFLICT')
    with zipfile.ZipFile(target) as z:
        media_hashes={__import__('hashlib').sha256(z.read(name)).hexdigest() for name in z.namelist() if name.startswith('ppt/media/')}
    if any(row['source_sha256'] not in media_hashes for row in checked['retained_assets']['assets']):raise GateError('SOURCE_IMAGE_BYTES_CHANGED')
    record_trace(out,contract,target,[dict(slide_id='SCENE',object_id=identity,native_id=shape.shape_id) for identity,shape in shapes.items()])
    verify_trace(out/'来源与引用链.json')
    ledger=from_scene(g)
    for row in ledger['elements']:
        if row['id'] in plans:row.update(treatment='native_bound_connector',editability='NATIVE_RECTANGLE_ENDPOINTS')
    for asset in checked['retained_assets']['assets']:
        ledger['required_ids'].append(asset['id']);ledger['elements'].append(dict(id=asset['id'],source_ref=asset['source_ref'],output_ids=[asset['id']],treatment='retained_pixels',editability='MOVE_SCALE_REPLACE',content_kind=asset['content_kind'],source_sha256=asset['source_sha256'],crop_box_normalized=asset['crop_box_normalized'],uncertain=False,review_status='PENDING'))
    add_unresolved(ledger,checked);save(out/'元素处理清单.json',ledger);save(out/'生产前检查.json',checked)
    unresolved=checked['reconstruction'].get('uncertain_ids',[])
    if unresolved:
        (out/'待核对事项.md').write_text('# 待核对事项\n\n以下源图元素未重建，保留待核对，不补猜文字：\n\n'+''.join('- '+identity+'\n' for identity in unresolved)+'\n坐标与来源哈希见元素处理清单和生产前检查。本文件不是正式图或作者接受证明。\n',encoding='utf-8')
    result=dict(status='GENERATED_RELOADED_PENDING_NATIVE_VISUAL_REVIEW',backend='python-pptx',file=source(target),scene=source(scene_path),
                generated_without_office=True,office_native_reopen='NOT_PERFORMED_BY_THIS_BACKEND',visual='PENDING',author_acceptance='PENDING',
                native_bound_connectors=plans,retained_assets=len(checked['retained_assets']['assets']),source_pixels_preserved=True,
                stable_ids='PresentationML Tags: STUDIO_SLIDE_ID / STUDIO_ID / ASSET_ID plus provenance contract hash',source_image=source(source_image) if source_image else None)
    result['unresolved_reconstruction_ids']=unresolved;result['release_ready']=False
    save(out/'生产验收.json',result);return result


def validate_supported(g):
    if g['asset_refs'] or g['clipping']:raise GateError('PORTABLE_GRAPH_ASSET_OR_CLIP_UNSUPPORTED_USE_EXPLICIT_RETAINED_ASSETS')
    if g['canvas']['unit']!='px' or g['canvas']['origin']!='TOP_LEFT':raise GateError('PORTABLE_COORDINATE_UNSUPPORTED')
    items=sorted([*g['objects'],*g['arrow_objects'],*g['text_objects']],key=lambda o:o['z_order'])
    if any(o['geometry']['kind'] not in ('RECT','ELLIPSE','TEXT','POLYLINE') for o in items):raise GateError('PORTABLE_GEOMETRY_UNSUPPORTED_USE_EXISTING_NATIVE_ROUTE')
    for obj in items:
        d=obj['geometry']['data'];style=g['style_tokens'][obj['style_token']]
        if style.get('opacity',1) not in (1,None):raise GateError('PORTABLE_OPACITY_UNSUPPORTED')
        if obj in g['arrow_objects'] and d.get('encoding') not in ('solid_arrow','dashed_arrow','line'):raise GateError('PORTABLE_ARROW_SEMANTICS_UNSUPPORTED')


def tag_scene(slide,shapes,scene_path,checked,slide_id):
    from provenance import scene_trace, object_tags
    from pptx_tags import attach_tags, shape_nv
    contract=scene_trace(scene_path,checked)
    attach_tags(slide,slide.element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}cSld'),dict(STUDIO_SLIDE_ID=slide_id))
    retained={a['id']:a for a in checked['retained_assets']['assets']}
    for identity,shape in shapes.items():
        tags=dict(STUDIO_ID=identity,**object_tags(contract,identity))
        if identity in retained:
            asset=retained[identity];tags.update(ASSET_ID=asset.get('asset_id',identity),ASSET_REVISION=asset.get('asset_revision','r1'),SEMANTIC_SHA256=contract['binding']['sha256'])
        attach_tags(slide,shape_nv(shape),tags)
    return contract


def draw_scene(prs, g, checked, scene_path):
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    plans=rectangle_connector_plan(g)
    scale=g['geometry']['coordinate_transform']['scale_pt_per_px']
    items=sorted([*g['objects'],*g['arrow_objects'],*g['text_objects']],key=lambda o:o['z_order'])
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
    shapes={};z_orders={o['id']:o['z_order'] for o in items}
    def color(value):return RGBColor.from_string(value.lstrip('#'))
    def metadata(shape,identity,details):
        shape.name=identity
        nv=shape.element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
        nv.set('descr',json.dumps(dict(baihe_sci_id=identity,scene_sha256=sha(scene_path),**details),ensure_ascii=False))
    def paint(shape,style):
        # PowerPoint's default theme otherwise adds shadows to new shapes/lines.
        sppr=shape.element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
        if sppr is not None:
            effects=sppr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
            if effects is None:sppr.append(OxmlElement('a:effectLst'))
        effect_ref=shape.element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef')
        if effect_ref is not None:effect_ref.set('idx','0')
        if hasattr(shape,'fill'):
            if style.get('fill') and style['fill']!='none':shape.fill.solid();shape.fill.fore_color.rgb=color(style['fill'])
            else:shape.fill.background()
        if style.get('stroke') and style.get('stroke_width',0)>0:
            shape.line.color.rgb=color(style['stroke']);shape.line.width=Pt(style['stroke_width']*scale)
        else:shape.line.fill.background()
    for obj in items:
        identity=obj['id'];d=obj['geometry']['data'];kind=obj['geometry']['kind'];style=g['style_tokens'][obj['style_token']]
        x0,y0,x1,y1=obj['bbox']
        if kind=='TEXT':
            size=d['font_size']*scale;text=d['text'];width=max((x1-x0)*scale,len(text)*size*.55)
            x=d['x']*scale-(width/2 if d['anchor']=='middle' else width if d['anchor']=='end' else 0)
            shape=slide.shapes.add_textbox(Pt(x),Pt((d['y']-d['font_size'])*scale),Pt(width),Pt(size*1.5))
            frame=shape.text_frame;frame.clear();frame.word_wrap=False;frame.vertical_anchor=MSO_ANCHOR.TOP
            frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
            paragraph=frame.paragraphs[0];paragraph.alignment={'start':PP_ALIGN.LEFT,'middle':PP_ALIGN.CENTER,'end':PP_ALIGN.RIGHT}[d['anchor']]
            paragraph.space_before=paragraph.space_after=Pt(0)
            run=paragraph.add_run();run.text=text;run.font.name=obj.get('font_family','Arial');run.font.size=Pt(size);run.font.color.rgb=color(style.get('fill') or '#263238')
            east_asian=OxmlElement('a:ea');east_asian.set('typeface','Microsoft YaHei');run._r.get_or_add_rPr().append(east_asian)
        elif kind in ('RECT','ELLIPSE'):
            shape=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE if kind=='RECT' else MSO_AUTO_SHAPE_TYPE.OVAL,Pt(d['x']*scale),Pt(d['y']*scale),Pt(d['width']*scale),Pt(d['height']*scale));paint(shape,style)
        else:
            points=d['points']
            if identity in plans:
                shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(points[0][0]*scale),Pt(points[0][1]*scale),Pt(points[-1][0]*scale),Pt(points[-1][1]*scale))
            else:
                builder=slide.shapes.build_freeform(points[0][0],points[0][1],scale=Pt(scale))
                builder.add_line_segments(points[1:],close=False);shape=builder.convert_to_shape()
            paint(shape,dict(style,fill=None))
            if d.get('dash') or d.get('encoding')=='dashed_arrow':shape.line.dash_style=MSO_LINE_DASH_STYLE.DASH
            if d.get('encoding') in ('solid_arrow','dashed_arrow'):
                end=OxmlElement('a:tailEnd');end.set('type','arrow');end.set('w','med');end.set('len','med');shape.line._get_or_add_ln().append(end)
        metadata(shape,identity,dict(scientific_refs=obj['scientific_refs'],source_layer_id=obj['layer_id']));shapes[identity]=shape
    for asset in checked['retained_assets']['assets']:
        x0,y0,x1,y1=asset['bbox'];shape=slide.shapes.add_picture(asset['source_path'],Pt(x0*scale),Pt(y0*scale),Pt((x1-x0)*scale),Pt((y1-y0)*scale))
        left,top,right,bottom=asset['crop_box_normalized'];shape.crop_left=left;shape.crop_top=top;shape.crop_right=1-right;shape.crop_bottom=1-bottom
        metadata(shape,asset['id'],dict(source_sha256=asset['source_sha256'],editability='MOVE_SCALE_REPLACE',crop_box_px=asset['source_box_px']))
        shapes[asset['id']]=shape;z_orders[asset['id']]=asset.get('z_order',0)
    links=list(g['parent_child'])+[dict(parent_id=a['group_id'],child_id=a['id']) for a in checked['retained_assets']['assets'] if a.get('group_id')]
    pending={group['id']:group for group in g['groups']}
    while pending:
        progress=False
        for identity,group in list(pending.items()):
            kids=[link['child_id'] for link in links if link['parent_id']==identity]
            if any(k in pending for k in kids):continue
            if not kids or any(k not in shapes for k in kids):raise GateError('PORTABLE_GROUP_UNRESOLVED')
            shape=slide.shapes.add_group_shape([shapes[k] for k in sorted(kids,key=lambda k:z_orders[k])]);metadata(shape,identity,dict(scientific_refs=group.get('scientific_refs',[])))
            shapes[identity]=shape;z_orders[identity]=min(z_orders[k] for k in kids);pending.pop(identity);progress=True
        if not progress:raise GateError('PORTABLE_PARENT_CYCLE')
    for identity,plan in plans.items():
        connector=shapes[identity];connector.begin_connect(shapes[plan['source']['object_id']],plan['source']['site']-1);connector.end_connect(shapes[plan['target']['object_id']],plan['target']['site']-1)
    for shape in sorted(list(slide.shapes),key=lambda s:z_orders[s.name]):slide.shapes._spTree.append(shape.element)
    return slide, shapes, plans


if __name__=='__main__':
    import sys
    print(json.dumps(produce(*sys.argv[1:]),ensure_ascii=False))
