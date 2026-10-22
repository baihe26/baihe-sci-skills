"""Behavioural regression cases for configuration, concurrent writers and scientific display contracts."""
from pathlib import Path
import sys, os, json, unittest, uuid, subprocess, shutil, copy
from unittest.mock import patch
REPO=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(REPO/'工作台/源码'))
from paths import ROOT, read, save, sha, digest
from core import GateError, new_task
from runtime_config import select_config
from preflight import inspect_pack
from refinement import revise
from project_intake import inspect_project, panel_candidates
from delivery_review import verify_delivery


class TargetedImprovements(unittest.TestCase):
    def temp(self):
        p=ROOT/'临时文件'/('定向升级_'+uuid.uuid4().hex);p.mkdir(parents=True);return p

    def test_config_priority_and_explicit_missing_do_not_switch(self):
        root=self.temp();personal=root/'baihe-sci.local.json';public=root/'baihe-sci.public.local.json'
        self.assertEqual(select_config(root,{})[1],'NOT_CONFIGURED')
        public.write_text('{}');self.assertEqual(select_config(root,{})[0],public)
        personal.write_text('{"keep":true}');self.assertEqual(select_config(root,{})[0],personal)
        self.assertEqual(select_config(root,{'BAIHE_SCI_CONFIG':str(public)})[0],public)
        missing=root/'not_there.json'
        self.assertEqual(select_config(root,{'BAIHE_SCI_CONFIG':str(missing)})[0],missing)
        moved=root/'moved';moved.mkdir();shutil.move(str(personal),str(moved/personal.name));shutil.move(str(public),str(moved/public.name))
        self.assertEqual(select_config(moved,{})[0],moved/personal.name)

    @unittest.skipUnless(os.name=='nt','Windows entrypoint comparison')
    def test_batch_and_direct_entrypoints_keep_personal_configuration(self):
        root=self.temp();shutil.copy2(REPO/'baihe-sci.cmd',root/'baihe-sci.cmd')
        shutil.copytree(REPO/'工作台/源码',root/'工作台/源码')
        env=dict(os.environ,BAIHE_SCI_PYTHON=sys.executable,PYTHONUTF8='1');env.pop('BAIHE_SCI_CONFIG',None)
        public=root/'baihe-sci.public.local.json';save(public,dict(portable_python=sys.executable))
        for personal in (False,True):
            if personal:save(root/'baihe-sci.local.json',dict(portable_python=sys.executable,keep='personal'))
            reports=[]
            for cmd,shell in (([sys.executable,str(root/'工作台/源码/workbench.py'),'--doctor'],False),(f'call "{root/"baihe-sci.cmd"}" --doctor',True)):
                result=subprocess.run(cmd,shell=shell,env=env,capture_output=True,text=True,encoding='utf-8',timeout=20)
                self.assertEqual(result.returncode,0,result.stdout+result.stderr);reports.append(json.loads(result.stdout))
            self.assertEqual(reports[0]['configuration'],reports[1]['configuration'])
            self.assertEqual(reports[0]['configuration']['selection'],'EXISTING_PERSONAL' if personal else 'PUBLIC_DISTRIBUTION')

    def test_cross_checkout_processes_share_one_writer_and_preserve_busy_marker(self):
        root=self.temp();shared=root/'共同.lock'
        code="""import sys
from pathlib import Path
sys.path.insert(0,sys.argv[1])
import coordination as c
c.ROOT=Path(sys.argv[2]);c.DESKTOP_LOCKS=[c.ROOT/'桌面生产.lock'];c.SHARED_WRITER_LOCKS=[Path(sys.argv[3])]
@c.single_writer
def run():
 print('ENTERED',flush=True)
 if sys.argv[4]=='hold':sys.stdin.readline()
try:run()
except Exception as e:print(str(e),flush=True);sys.exit(2)
"""
        args=[sys.executable,'-c',code,str(REPO/'工作台/源码')]
        first=subprocess.Popen([*args,str(root/'one'),str(shared),'hold'],stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,encoding='utf-8')
        try:
            self.assertEqual(first.stdout.readline().strip(),'ENTERED')
            second=subprocess.run([*args,str(root/'two'),str(shared),'once'],capture_output=True,text=True,encoding='utf-8',timeout=15)
            self.assertEqual(second.returncode,2);self.assertNotIn('ENTERED',second.stdout)
            self.assertIn('WRITER_BUSY',second.stdout)
        finally:
            first.communicate('\n',timeout=15)
        self.assertEqual(first.returncode,0);self.assertFalse(shared.exists())
        shared.write_text('UNKNOWN_NATIVE_OUTCOME',encoding='utf-8')
        third=subprocess.run([*args,str(root/'three'),str(shared),'once'],capture_output=True,text=True,encoding='utf-8',timeout=15)
        self.assertEqual(third.returncode,2);self.assertEqual(shared.read_text(),'UNKNOWN_NATIVE_OUTCOME')

    def request(self,folder,scene,decorative=False,target='AI_text_A'):
        review=dict(artifact_scene_sha256=sha(scene),reviewer='SYNTHETIC_REVIEW',findings=[dict(id='F1',object_ids=[target])])
        if decorative:review['decorative_color_permissions']=[dict(object_id=target,fields=['fill'],reason='Explicit synthetic decorative contrast review')]
        save(folder/'review.json',review)
        request=dict(expected_parent_sha256=sha(scene),expected_pack_sha256=inspect_pack(scene)['input_fingerprint'],
            review=dict(path='review.json',sha256=sha(folder/'review.json')),finding_ids=['F1'],allowed_object_ids=[target],
            changes=[dict(object_id=target,fields=dict(fill='#102030'))])
        save(folder/'request.json',request);return folder/'request.json'

    def test_unknown_colour_is_not_automatically_scientifically_equivalent(self):
        from color_roles import permit_display_color
        self.assertEqual(permit_display_color('A','fill','#aabbcc','#AABBCC',{},{}),'UNCHANGED')
        scene=REPO/'示例/混合重建/IllustratorSceneGraph.json';folder=self.temp()
        request=self.request(folder,scene)
        with self.assertRaisesRegex(GateError,'COLOR_MEANING_UNKNOWN'):revise(scene,request,folder/'blocked')
        self.assertFalse((folder/'blocked').exists())
        folder=self.temp();request=self.request(folder,scene,True)
        result=revise(scene,request,folder/'allowed')
        self.assertEqual(result['color_changes'][0]['permission'],'REVIEWED_DECORATIVE')

    def test_scientific_colour_cannot_be_downgraded_by_generic_review(self):
        folder=self.temp();base=folder/'pack';shutil.copytree(REPO/'示例/混合重建',base);scene=base/'IllustratorSceneGraph.json'
        graph=read(scene);obj=next(o for o in graph['objects'] if o['id']=='AI_box_A');value=graph['style_tokens'][obj['style_token']]['fill']
        save(base/'color_roles.json',dict(schema_version='color-roles-0.1',bindings=[dict(object_id='AI_box_A',field='fill',kind='scientific',role='material_identity',value=value,reason='Synthetic group encoding')]))
        request=self.request(folder,scene,True,'AI_box_A')
        with self.assertRaisesRegex(GateError,'SCIENTIFIC_COLOR_REENCODING'):revise(scene,request,folder/'blocked')
        roles=read(base/'color_roles.json');roles['bindings'][0]['value']='#FFFFFF';(base/'color_roles.json').write_text(json.dumps(roles),encoding='utf-8')
        with self.assertRaisesRegex(GateError,'COLOR_ROLE_CONFLICT'):inspect_pack(scene)

    def test_style_divergence_and_sidecar_changes_invalidate_input_pack(self):
        folder=self.temp();base=folder/'pack';shutil.copytree(REPO/'示例/混合重建',base);scene=base/'IllustratorSceneGraph.json'
        save(base/'color_roles.json',dict(schema_version='color-roles-0.1',bindings=[dict(object_id='AI_text_A',field='fill',kind='decorative',reason='Synthetic label')]))
        request=self.request(folder,scene)
        roles=read(base/'color_roles.json');roles['bindings'][0]['reason']='Changed classification evidence';(base/'color_roles.json').write_text(json.dumps(roles))
        with self.assertRaisesRegex(GateError,'INPUT_PACK_CONFLICT'):revise(scene,request,folder/'blocked')
        graph=read(scene);obj=graph['objects'][0];token=obj['style_token'];graph['style_tokens'][token]['fill']='#FFFFFF'
        scene.write_text(json.dumps(graph),encoding='utf-8')
        with self.assertRaisesRegex(GateError,'STYLE_CONFLICT'):inspect_pack(scene)

    def project(self,folder):
        from PIL import Image
        for name,size in [('a.png',(800,400)),('b.png',(1000,500))]:Image.new('RGB',size,'white').save(folder/name)
        m=dict(schema_version='project-intake-0.1',project_id='SYNTHETIC_PANEL_TEST',canvas_mm=[180,120],layout_columns=[1,2],
            sources=[dict(id=k,path=n,kind='data_plot',review_status='CONFIRMED') for k,n in [('A','a.png'),('B','b.png')]],
            panels=[dict(id='a',source_id='A',plot_area_normalized=[.1,.1,.9,.85]),dict(id='b',source_id='B',plot_area_normalized=[.2,.2,.95,.9])],
            required_figure_ids=['Fig1','Fig2'],figure_uses=[dict(source_figure_id='Fig1',placements=['panel_a'])])
        save(folder/'project.json',m);return folder/'project.json'

    def test_panel_plan_aligns_content_without_stretch_and_preserves_source(self):
        folder=self.temp();path=self.project(folder);before=[sha(folder/n) for n in ('a.png','b.png')]
        result=panel_candidates(path,folder/'out');self.assertEqual(len(result['candidates']),2)
        panels=result['candidates'][1]['panels'];self.assertAlmostEqual(panels[0]['plot_area_mm'][3],panels[1]['plot_area_mm'][3])
        for row in panels:
            x,y,r,b=row['bbox_mm'];self.assertAlmostEqual((r-x)/(b-y),2)
            self.assertEqual(row['font_check']['status'],'NOT_RECORDED_SOURCE_TEXT_METRICS_REQUIRE_REVIEW')
        self.assertEqual(before,[sha(folder/n) for n in ('a.png','b.png')])
        self.assertEqual(result['intake']['unaccounted_figure_ids'],['Fig2']);self.assertFalse(result['production_started'])

    def test_intake_unknown_facts_and_shared_crop_conflicts_are_not_guessed(self):
        folder=self.temp();path=self.project(folder);m=read(path)
        m['sources'].append(dict(id='raw',path='a.png',kind='numerical_data'))
        for panel in m['panels']:panel.update(crop_group='same_field',frame_id='F1')
        path.write_text(json.dumps(m));report=inspect_project(path)
        self.assertIn('statistical_n',report['pending'][0]['missing_fields'])
        with self.assertRaisesRegex(GateError,'SHARED_CROP_CONFLICT'):panel_candidates(path,folder/'blocked')
        self.assertFalse((folder/'blocked').exists())

    def test_existing_plot_explicit_unknown_remains_pending_even_if_source_confirmed(self):
        folder=self.temp();path=self.project(folder);m=read(path)
        m['sources'][0]['scientific_metadata']=dict(statistical_n='UNKNOWN',nested=dict(denominator='待核对'))
        path.write_text(json.dumps(m),encoding='utf-8')
        from project_intake import intake
        report=intake(path,folder/'intake')
        self.assertIn('statistical_n',report['pending'][0]['missing_fields'])
        self.assertIn('nested.denominator',report['pending'][0]['missing_fields'])
        self.assertNotIn('experimental_unit',report['pending'][0]['missing_fields'])
        self.assertIn('statistical_n',(folder/'intake/从这里开始.md').read_text(encoding='utf-8'))
        self.assertEqual(report['sources'][0]['review_status'],'CONFIRMED')

    def test_layout_threshold_is_positive_and_zero_margin_remains_available(self):
        folder=self.temp();path=self.project(folder);m=read(path)
        m['minimum_font_pt']=-1;path.write_text(json.dumps(m))
        with self.assertRaisesRegex(GateError,'MINIMUM_FONT_INVALID'):panel_candidates(path,folder/'bad')
        m.update(minimum_font_pt=7,margin_mm=0,gap_mm=0,label_band_mm=0);path.write_text(json.dumps(m))
        report=panel_candidates(path,folder/'valid')
        self.assertEqual(report['status'],'PANEL_LAYOUT_CANDIDATES_NOT_PRODUCTION')

    def test_final_object_check_rejects_changed_text_and_keeps_edit_test_pending(self):
        from portable_ppt import produce
        from pptx import Presentation
        folder=self.temp();produced=produce(REPO/'示例/混合重建/IllustratorSceneGraph.json',folder/'ppt');artifact=Path(produced['file']['path'])
        goals=dict(schema_version='delivery-requirements-0.1',observations=[dict(id='T1',object_id='AI_text_A',kind='text_equals',expected='Module A'),dict(id='E1',object_id='AI_box_A',kind='edit_action',expected='move_group')])
        save(folder/'goals.json',goals)
        self.assertEqual(verify_delivery(folder/'goals.json',artifact)['status'],'PENDING_NATIVE_CHECKS')
        prs=Presentation(artifact)
        def walk(shapes):
            for s in shapes:
                if s.name=='AI_text_A':s.text='Module B'
                if s.shape_type==6:walk(s.shapes)
        walk(prs.slides[0].shapes);prs.save(folder/'changed.pptx')
        result=verify_delivery(folder/'goals.json',folder/'changed.pptx')
        self.assertEqual(result['status'],'FAIL');self.assertEqual(result['artifact']['sha256'],sha(folder/'changed.pptx'))

    def test_profile_sources_and_overrides_are_bound_and_never_inherit_acceptance(self):
        folder=self.temp();base=folder/'pack';shutil.copytree(REPO/'示例/混合重建',base);scene=base/'IllustratorSceneGraph.json'
        save(base/'synthetic_style.json',dict(font='synthetic',author_acceptance='PENDING'))
        save(base/'style_profile.json',dict(schema_version='style-profile-0.1',profile_id='SYNTHETIC',version='1',scope='THIS_REVISION',source=dict(path='synthetic_style.json',sha256=sha(base/'synthetic_style.json'))))
        check=inspect_pack(scene);self.assertEqual(check['style_profile']['author_acceptance'],'NOT_INHERITED')
        (base/'synthetic_style.json').write_text('{}')
        with self.assertRaisesRegex(GateError,'SOURCE_HASH_CONFLICT'):inspect_pack(scene)

    def test_preflight_failure_does_not_allocate_request_and_unknown_journal_wins(self):
        from task_status import status
        from recovery_journal import Journal
        folder=self.temp();key='提前预检_'+uuid.uuid4().hex
        env=dict(os.environ,BAIHE_SCI_CONFIG=str(folder/'absent.json'),PYTHONUTF8='1')
        result=subprocess.run([sys.executable,str(REPO/'工作台/源码/workbench.py'),'--module','vector','--scene',str(REPO/'示例/简单科学关系/IllustratorSceneGraph.json'),'--request-id',key],env=env,capture_output=True,text=True,encoding='utf-8',timeout=15)
        self.assertEqual(result.returncode,2);self.assertFalse((ROOT/'任务'/key).exists())
        state=new_task('只作图','SYNTHETIC',[],key);task=ROOT/'任务'/key
        save(task/'执行回执.json',dict(output_manifest=[]));Journal(task/'恢复日志.jsonl',digest(state)).begin(None,dict(id='synthetic_native'))
        self.assertEqual(status(key)['status'],'OUTCOME_UNKNOWN_RECONCILE_FIRST')


if __name__=='__main__':unittest.main()
