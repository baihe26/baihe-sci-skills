"""Saved-file interoperability and scientific/conflict regression, using synthetic packs."""
from pathlib import Path
import sys, unittest, uuid, copy, shutil, json, zipfile, os, subprocess, hashlib
REPO=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(REPO/'工作台/源码'))
from paths import ROOT,read,save,sha,digest
from core import GateError
from provenance import scene_trace,verify_trace
from portable_ppt import produce
from portable_deck import produce_deck
from pptx_tags import presentation_locations,read_package,xml,serialized,NS
from reference_update import index_references,propagate,propagate_request


class LineageUpdateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base=ROOT/'临时文件'/('链路测试_'+uuid.uuid4().hex);cls.base.mkdir(parents=True)
        cls.scene=REPO/'示例/混合重建/IllustratorSceneGraph.json'
        cls.original=cls.base/'original';produce(cls.scene,cls.original)
        cls.deck=cls.original/'可编辑示意图.pptx'
        from PIL import Image
        cls.asset=cls.base/'新版合成纹理.png';Image.new('RGB',(900,500),'#247268').save(cls.asset)
        cls.semantic=cls.scene.parent/'scientific_binding.json'

    def temp(self):
        p=self.base/uuid.uuid4().hex;p.mkdir();return p

    def clone(self,folder=None):
        folder=folder or self.temp();path=folder/'same.pptx';shutil.copy2(self.deck,path);return path

    def changed_package(self,path,change):
        parts=read_package(path);change(parts)
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for name,blob in parts.items():z.writestr(name,blob)

    def test_claim_evidence_object_chain_does_not_infer_causality(self):
        contract=scene_trace(self.scene);objects={r['object_id']:r for r in contract['objects']}
        self.assertEqual(objects['AI_R1']['direct_claim_ids'],['C1'])
        self.assertEqual(objects['AI_R2']['direct_claim_ids'],['C2'])
        self.assertEqual(objects['AI_box_A']['direct_claim_ids'],[])
        self.assertEqual(objects['AI_box_A']['contextual_claim_ids'],['C1','C2'])
        self.assertEqual(objects['texture_inset']['direct_claim_ids'],[])
        claims={r['claim']['id']:r for r in contract['claims']}
        self.assertEqual(claims['C2']['claim']['evidence_level'],'HYPOTHESIZED')
        self.assertEqual(claims['C2']['evidence_links'][0]['relation'],'limits')
        self.assertFalse(contract['automatic_adoption'])

    def test_real_native_tags_and_complete_trace_readback(self):
        rows=presentation_locations(self.deck);self.assertEqual(len(rows),13)
        self.assertEqual({r['slide_id'] for r in rows},{'SCENE'})
        self.assertTrue(all(r['tags']['PROVENANCE_SHA256'] for r in rows))
        self.assertEqual(verify_trace(self.original/'来源与引用链.json')['occurrences'],13)
        refs=index_references(self.deck);self.assertEqual(len(refs),1)
        self.assertEqual(refs[0]['asset_id'],'texture_inset')
        self.assertEqual(refs[0]['semantic_sha256'],sha(self.semantic))

    def test_trace_cannot_drop_occurrences_or_silently_rebind_artifact(self):
        folder=self.temp();trace=read(self.original/'来源与引用链.json');trace['occurrences'].pop();save(folder/'drop.json',trace)
        with self.assertRaisesRegex(GateError,'COVERAGE'):verify_trace(folder/'drop.json')
        deck=self.clone();deck.write_bytes(deck.read_bytes()+b'edited')
        with self.assertRaisesRegex(GateError,'SOURCE_HASH_CONFLICT'):verify_trace(self.original/'来源与引用链.json',deck)

    def test_trace_detects_changed_sidecar(self):
        folder=self.temp();pack=folder/'pack';shutil.copytree(self.scene.parent,pack)
        produce(pack/'IllustratorSceneGraph.json',folder/'new')
        evidence=read(pack/'evidence_map.json');evidence['decisions']=[dict(id='new',status='PROPOSED')]
        (pack/'evidence_map.json').write_text(json.dumps(evidence),encoding='utf-8')
        with self.assertRaisesRegex(GateError,'INPUT_PACK_CHANGED'):verify_trace(folder/'new/来源与引用链.json')

    def test_multiple_pages_keep_duplicate_object_names_scoped_by_slide(self):
        folder=self.temp();manifest=dict(schema_version='scene-deck-0.1',slides=[dict(id=k,scene_path=str(self.scene),expected_scene_sha256=sha(self.scene)) for k in ('ONE','TWO')]);save(folder/'deck.json',manifest)
        report=produce_deck(folder/'deck.json',folder/'deck')
        loc=presentation_locations(report['file']['path']);self.assertEqual(len(loc),26)
        refs=index_references(report['file']['path']);self.assertEqual(len(refs),2)
        self.assertEqual({r['slide_id'] for r in refs},{'ONE','TWO'})
        self.assertEqual(len({r['reference_id'] for r in refs}),2)
        self.assertTrue(all(p['verified']['occurrences']==13 for p in report['pages']))
        self.assertTrue((folder/'deck/待核对事项.md').is_file())

    def test_duplicate_slide_ids_and_wrong_hash_fail_before_output(self):
        for bad in ('duplicate','hash'):
            folder=self.temp();rows=[dict(id=k,scene_path=str(self.scene),expected_scene_sha256=sha(self.scene)) for k in ('ONE','TWO')]
            if bad=='duplicate':rows[1]['id']='ONE'
            else:rows[1]['expected_scene_sha256']='0'*64
            save(folder/'deck.json',dict(schema_version='scene-deck-0.1',slides=rows))
            with self.subTest(bad=bad),self.assertRaises(GateError):produce_deck(folder/'deck.json',folder/'blocked')
            self.assertFalse((folder/'blocked').exists())

    def test_surgical_copy_update_preserves_unrelated_parts_and_original_media(self):
        deck=self.clone();old=sha(deck);before=read_package(deck);refs=index_references(deck);folder=self.temp()
        result=propagate(refs,self.asset,'r2',folder/'updated',self.semantic)
        self.assertEqual(sha(deck),old);self.assertEqual(len(result['results'][0]['updated']),1)
        target=result['results'][0]['output']['path'];after=read_package(target)
        for name in before:
            if name not in result['results'][0]['changed_original_parts']:self.assertEqual(before[name],after[name])
            if name.startswith('ppt/media/'):self.assertEqual(before[name],after[name])
        new=index_references(target)[0];self.assertEqual(new['media_sha256'],sha(self.asset));self.assertEqual(new['revision'],'r2')
        loc=next(r for r in presentation_locations(target) if r['object_id']=='texture_inset')
        self.assertNotIn('PROVENANCE_SHA256',loc['tags'])
        self.assertEqual(loc['tags']['PROVENANCE_STATUS'],'REPLACEMENT_PIXELS_PENDING_SOURCE_REVIEW')
        self.assertTrue(loc['tags']['PREVIOUS_PROVENANCE_SHA256'])
        self.assertEqual(refs[0]['media_sha256'],sha(self.scene.parent/'合成原图.png'))

    def test_same_named_decks_and_manual_conflict_preserved(self):
        a=self.clone();b=self.clone();refs=index_references(a)+index_references(b)
        self.assertNotEqual(refs[0]['reference_id'],refs[1]['reference_id'])
        from pptx import Presentation
        p=Presentation(b);picture=next(s for s in p.slides[0].shapes if s.name=='texture_inset');picture.left+=100000;p.save(b)
        source_hashes=[sha(a),sha(b)]
        result=propagate(refs,self.asset,'r2',self.temp()/'updates',self.semantic)
        self.assertEqual(len(result['results'][0]['updated']),1)
        self.assertEqual(len(result['results'][1]['conflicts']),1)
        self.assertIsNone(result['results'][1]['output'])
        self.assertEqual(source_hashes,[sha(a),sha(b)])

    def test_group_transform_is_part_of_reference_fingerprint(self):
        # Place the retained image in an existing semantic group in a new pack.
        folder=self.temp();pack=folder/'pack';shutil.copytree(self.scene.parent,pack)
        assets=read(pack/'retained_assets.json');assets['assets'][0]['group_id']='AI_fig_A'
        (pack/'retained_assets.json').write_text(json.dumps(assets),encoding='utf-8')
        produce(pack/'IllustratorSceneGraph.json',folder/'deck');deck=folder/'deck/可编辑示意图.pptx';refs=index_references(deck)
        from pptx import Presentation
        p=Presentation(deck);next(s for s in p.slides[0].shapes if s.name=='AI_fig_A').left+=100000;p.save(deck)
        result=propagate(refs,self.asset,'r2',folder/'update',pack/'scientific_binding.json')
        self.assertEqual(len(result['results'][0]['conflicts']),1)
        self.assertIsNone(result['results'][0]['output'])

    def test_explicit_override_wrong_semantics_and_duplicate_targets(self):
        deck=self.clone();refs=index_references(deck)
        folder=self.temp();wrong=folder/'wrong.json';wrong.write_text('{}',encoding='utf-8')
        with self.assertRaisesRegex(GateError,'SEMANTIC'):propagate(refs,self.asset,'r2',folder/'bad_semantics',wrong)
        with self.assertRaisesRegex(GateError,'DUPLICATE'):propagate(refs+refs,self.asset,'r2',folder/'duplicate',self.semantic)
        refs[0]['manual_override']=True
        r=propagate(refs,self.asset,'r2',folder/'conflict',self.semantic);self.assertIsNone(r['results'][0]['output'])
        self.assertFalse((folder/'bad_semantics').exists());self.assertFalse((folder/'duplicate').exists())

    def test_different_aspect_ratio_rejected_without_stretch(self):
        from PIL import Image
        folder=self.temp();asset=folder/'square.png';Image.new('RGB',(100,100),'white').save(asset)
        with self.assertRaisesRegex(GateError,'ASPECT_CHANGE'):propagate(index_references(self.deck),asset,'r2',folder/'bad',self.semantic)
        self.assertFalse((folder/'bad').exists())

    def test_hash_bound_request_and_subset_selection(self):
        deck=self.clone();folder=self.temp();refs=index_references(deck);save(folder/'index.json',dict(references=refs))
        request=dict(reference_index=dict(path='index.json',sha256=sha(folder/'index.json')),reference_ids=[refs[0]['reference_id']],
                     new_asset=dict(path=str(self.asset),sha256=sha(self.asset)),new_revision='r2',semantic_contract=dict(path=str(self.semantic),sha256=sha(self.semantic)))
        save(folder/'request.json',request);r=propagate_request(folder/'request.json',folder/'good');self.assertTrue(r['source_files_unchanged'])
        request['new_asset']['sha256']='0'*64;save(folder/'bad.json',request)
        with self.assertRaisesRegex(GateError,'HASH_CONFLICT'):propagate_request(folder/'bad.json',folder/'bad')

    def test_native_duplicate_identity_is_not_guessed(self):
        deck=self.clone()
        def change(parts):
            tags=[k for k,v in parts.items() if k.startswith('ppt/tags/') and b'name="STUDIO_ID" val="AI_R1"' in v]
            self.assertEqual(len(tags),1);parts[tags[0]]=parts[tags[0]].replace(b'val="AI_R1"',b'val="texture_inset"')
        self.changed_package(deck,change)
        with self.assertRaisesRegex(GateError,'AMBIGUOUS'):index_references(deck)

    def test_slide_reordering_keeps_reference_identity(self):
        folder=self.temp();save(folder/'deck.json',dict(schema_version='scene-deck-0.1',slides=[dict(id=k,scene_path=str(self.scene),expected_scene_sha256=sha(self.scene)) for k in ('ONE','TWO')]))
        report=produce_deck(folder/'deck.json',folder/'deck');deck=Path(report['file']['path']);refs=index_references(deck)
        from pptx import Presentation
        p=Presentation(deck);first=p.slides._sldIdLst[0];p.slides._sldIdLst.remove(first);p.slides._sldIdLst.append(first);p.save(deck)
        result=propagate(refs,self.asset,'r2',folder/'updates',self.semantic)
        self.assertEqual(len(result['results'][0]['updated']),2);self.assertEqual(result['results'][0]['conflicts'],[])

    def test_shared_media_only_selected_reference_is_updated(self):
        folder=self.temp();save(folder/'deck.json',dict(schema_version='scene-deck-0.1',slides=[dict(id=k,scene_path=str(self.scene),expected_scene_sha256=sha(self.scene)) for k in ('ONE','TWO')]))
        result=produce_deck(folder/'deck.json',folder/'deck');deck=Path(result['file']['path']);refs=index_references(deck)
        self.assertEqual(refs[0]['media_part'],refs[1]['media_part'])
        r=propagate(refs[:1],self.asset,'r2',folder/'update',self.semantic)
        after=index_references(r['results'][0]['output']['path']);by_id={q['slide_id']:q for q in after}
        self.assertEqual(by_id['ONE']['media_sha256'],sha(self.asset))
        self.assertEqual(by_id['TWO']['object_digest'],refs[1]['object_digest'])
        self.assertEqual(by_id['TWO']['revision'],'r1')

    def test_connector_nested_owner_collinear_and_backtracking(self):
        from native_connectors import rectangle_connector_plan
        graph=read(self.scene);graph['arrow_objects'][0]['geometry']['data']['points']=[[300,220],[450,220],[600,220]]
        self.assertIn('AI_R1',rectangle_connector_plan(graph))
        graph['arrow_objects'][0]['geometry']['data']['points']=[[300,220],[500,220],[400,220],[600,220]]
        self.assertNotIn('AI_R1',rectangle_connector_plan(graph))
        graph['arrow_objects'][0]['geometry']['data']['points']=[[300,220],[600,220]]
        graph['parent_child']=[r for r in graph['parent_child'] if r['child_id']!='AI_box_A']+[dict(parent_id='AI_fig_A',child_id='INNER'),dict(parent_id='INNER',child_id='AI_box_A')]
        self.assertIn('AI_R1',rectangle_connector_plan(graph));self.assertNotIn('AI_R2',rectangle_connector_plan(graph))

    def test_multiscene_cli_records_complete_input_dependencies(self):
        folder=self.temp();key='链路CLI_'+uuid.uuid4().hex
        env=dict(os.environ,BAIHE_SCI_CONFIG=str(folder/'absent.json'),PYTHONUTF8='1')
        result=subprocess.run([sys.executable,str(REPO/'工作台/源码/workbench.py'),'--module','vector','--target','ppt','--backend','portable','--scene-deck',str(REPO/'示例/多页引用/scene_deck.json'),'--out',str(folder/'cli'),'--request-id',key],cwd=REPO,env=env,capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(result.returncode,0,result.stdout+result.stderr)
        task=read(ROOT/'任务'/key/'任务记录.json');self.assertTrue(any(Path(s['path']).name=='evidence_map.json' for s in task['sources']))
        self.assertEqual(len(list((ROOT/'任务'/key/'交接').glob('交接_*.json'))),2)


if __name__=='__main__':unittest.main()
